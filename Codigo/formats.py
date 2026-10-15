"""Extracción de texto por formato: PDF (con OCR opcional), DOCX, XLSX, PPTX, TXT y CSV."""
from __future__ import annotations

import csv
import io
import logging
import os
from pathlib import Path
from typing import Dict, List

import pdfplumber

import rag_core

LOGGER = logging.getLogger(__name__)
SUPPORTED = {".pdf", ".docx", ".xlsx", ".pptx", ".txt", ".csv"}
FILAS_POR_PAGINA = 30
LINEAS_POR_PAGINA = 40

_ocr_estado = None
_ocr_idioma = None


def ocr_disponible() -> bool:
    """True solo si están pytesseract, pypdfium2 y el programa Tesseract instalado."""
    global _ocr_estado
    if _ocr_estado is None:
        try:
            import pypdfium2  # noqa: F401
            import pytesseract
            if os.environ.get("VEXA_TESSERACT_CMD"):
                pytesseract.pytesseract.tesseract_cmd = os.environ["VEXA_TESSERACT_CMD"]
            pytesseract.get_tesseract_version()
            _ocr_estado = True
        except Exception:
            _ocr_estado = False
    return _ocr_estado


def _idioma_ocr() -> str:
    global _ocr_idioma
    if _ocr_idioma is None:
        import pytesseract
        langs = set(pytesseract.get_languages())
        _ocr_idioma = "+".join(l for l in ("spa", "eng") if l in langs) or "eng"
    return _ocr_idioma


def _ocr_pagina(datos: bytes, indice0: int) -> str:
    import pypdfium2 as pdfium
    import pytesseract
    imagen = pdfium.PdfDocument(datos)[indice0].render(scale=2.5).to_pil()
    return pytesseract.image_to_string(imagen, lang=_idioma_ocr())


def _pdf(datos: bytes) -> List[Dict]:
    paginas = []
    with pdfplumber.open(io.BytesIO(datos)) as pdf:
        for i, pagina in enumerate(pdf.pages, start=1):
            texto = (pagina.extract_text() or "").strip()
            if not texto and ocr_disponible():
                try:
                    texto = _ocr_pagina(datos, i - 1).strip()
                except Exception:
                    LOGGER.warning("OCR falló en la página %s", i)
            if texto:
                paginas.append({"page_number": i, "text": texto})
    return paginas


def _decodificar(datos: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return datos.decode(enc)
        except UnicodeDecodeError:
            continue
    return datos.decode("latin-1", errors="replace")


def _txt(datos: bytes) -> List[Dict]:
    texto = "\n".join(l.rstrip() for l in _decodificar(datos).splitlines())
    parrafos = [p for p in texto.split("\n\n") if p.strip()]
    paginas, grupo, n = [], [], 0
    for p in parrafos:
        grupo.append(p)
        n += p.count("\n") + 1
        if n >= LINEAS_POR_PAGINA:
            paginas.append({"page_number": len(paginas) + 1, "text": "\n\n".join(grupo)}); grupo, n = [], 0
    if grupo:
        paginas.append({"page_number": len(paginas) + 1, "text": "\n\n".join(grupo)})
    return paginas


def _filas_a_paginas(filas: List[List[str]], prefijo: str, pagina0: int = 1) -> List[Dict]:
    """Primera fila = encabezados. Cada fila pasa a 'Col: valor | Col: valor' y es su propio párrafo."""
    if not filas:
        return []
    cab = [c.strip() or f"Col{i + 1}" for i, c in enumerate(filas[0])]
    registros = []
    for fila in filas[1:]:
        celdas = [f"{cab[i] if i < len(cab) else f'Col{i + 1}'}: {v.strip()}"
                  for i, v in enumerate(fila) if v and v.strip()]
        if celdas:
            registros.append((prefijo + " | " if prefijo else "") + " | ".join(celdas))
    return [{"page_number": pagina0 + i // FILAS_POR_PAGINA,
             "text": "\n\n".join(registros[i:i + FILAS_POR_PAGINA])}
            for i in range(0, len(registros), FILAS_POR_PAGINA)]


def _csv(datos: bytes) -> List[Dict]:
    texto = _decodificar(datos)
    try:
        dialecto = csv.Sniffer().sniff(texto[:4096], delimiters=",;\t|")
    except csv.Error:
        dialecto = csv.excel
    return _filas_a_paginas(list(csv.reader(io.StringIO(texto), dialecto)), "")


def _xlsx(datos: bytes) -> List[Dict]:
    from openpyxl import load_workbook
    libro = load_workbook(io.BytesIO(datos), read_only=True, data_only=True)
    paginas = []
    for hoja in libro.worksheets:
        filas = [["" if v is None else str(v) for v in fila] for fila in hoja.iter_rows(values_only=True)]
        for p in _filas_a_paginas(filas, f"Hoja {hoja.title}", pagina0=len(paginas) + 1):
            p["page_number"] = len(paginas) + 1
            paginas.append(p)
    return paginas


def _pptx(datos: bytes) -> List[Dict]:
    from pptx import Presentation
    paginas = []
    for n, diapositiva in enumerate(Presentation(io.BytesIO(datos)).slides, start=1):
        textos = []
        for forma in diapositiva.shapes:
            if forma.has_text_frame and forma.text_frame.text.strip():
                textos.append(forma.text_frame.text.strip())
            if getattr(forma, "has_table", False) and forma.has_table:
                for fila in forma.table.rows:
                    celdas = [c.text.strip() for c in fila.cells if c.text.strip()]
                    if celdas:
                        textos.append(" | ".join(celdas))
        if diapositiva.has_notes_slide and diapositiva.notes_slide.notes_text_frame.text.strip():
            textos.append("Notas: " + diapositiva.notes_slide.notes_text_frame.text.strip())
        if textos:
            paginas.append({"page_number": n, "text": "\n\n".join(textos)})
    return paginas


def extraer(nombre: str, datos: bytes) -> List[Dict]:
    """Devuelve [{'page_number', 'text'}] según la extensión de `nombre`."""
    ext = Path(nombre).suffix.lower()
    if ext == ".pdf":
        return _pdf(datos)
    if ext == ".docx":
        return rag_core.extraer_texto_docx(io.BytesIO(datos))
    if ext == ".xlsx":
        return _xlsx(datos)
    if ext == ".pptx":
        return _pptx(datos)
    if ext == ".csv":
        return _csv(datos)
    if ext == ".txt":
        return _txt(datos)
    raise ValueError(f"Formato no soportado: {ext}")
