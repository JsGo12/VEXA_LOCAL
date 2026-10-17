import io, os, sys, csv
import numpy as np, pytest
sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))
import formats, vault, model_select, rag_stages, rag_hybrid, rag_core, store, watcher
from rag_core import Chunk


# ---------------- formatos
def test_txt_csv_xlsx_pptx(tmp_path):
    assert "Hola" in formats.extraer("a.txt", "Hola mundo\n\nSegundo párrafo".encode())[0]["text"]
    cuerpo = "Nombre;Teléfono\nAna;555-1\nLuis;555-2\n"
    t = formats.extraer("c.csv", cuerpo.encode("latin-1"))[0]["text"]
    assert "Nombre: Ana | Teléfono: 555-1" in t and "\n\n" in t
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = "Clientes"
    ws.append(["Nombre", "Monto"]); ws.append(["Ana", 1000])
    b = io.BytesIO(); wb.save(b)
    assert "Hoja Clientes | Nombre: Ana | Monto: 1000" in formats.extraer("x.xlsx", b.getvalue())[0]["text"]
    from pptx import Presentation
    p = Presentation(); s = p.slides.add_slide(p.slide_layouts[1])
    s.shapes.title.text = "Plan"; s.placeholders[1].text = "Entrega en marzo"
    b = io.BytesIO(); p.save(b)
    pag = formats.extraer("d.pptx", b.getvalue())
    assert pag[0]["page_number"] == 1 and "Entrega en marzo" in pag[0]["text"]
    with pytest.raises(ValueError):
        formats.extraer("x.exe", b"")


def test_ocr_real_si_hay_tesseract():
    if not formats.ocr_disponible():
        pytest.skip("Tesseract no instalado")
    from PIL import Image, ImageDraw, ImageFont
    img = Image.new("RGB", (900, 300), "white"); d = ImageDraw.Draw(img)
    try:
        f = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 48)
    except OSError:
        pytest.skip("sin fuente")
    d.text((30, 100), "Contrato de arriendo", fill="black", font=f)
    b = io.BytesIO(); img.save(b, "PDF")
    pag = formats.extraer("scan.pdf", b.getvalue())
    assert pag and "arriendo" in pag[0]["text"].lower()


# ---------------- cifrado
def test_vault_roundtrip_y_lectura_transparente(tmp_path):
    f = tmp_path / "a.bin"; f.write_bytes(b"secreto")
    assert vault.read_bytes(f) == b"secreto"                  # archivo antiguo sin cifrar
    assert vault.encrypt_existing(tmp_path) == 1
    assert b"secreto" not in f.read_bytes() and vault.read_bytes(f) == b"secreto"
    assert vault.encrypt_existing(tmp_path) == 0
    assert vault.dec_text(vault.enc_text("ñandú")) == "ñandú" and vault.dec_text("plano") == "plano"


def test_store_guarda_cifrado(tmp_path):
    db = tmp_path / "t.db"; store.init(db)
    store.save_chunks(db, "sha", "v2", [(0, 1, 1, "texto confidencial")])
    vectores = np.array([[0.125, 0.25, 0.5]], dtype=np.float32)
    store.save_vectors(db, "sha", "v2", "emb", vectores)
    store.add_message(db, 1, 1, "user", "pregunta privada")
    raw = open(db, "rb").read()
    import glob
    raw += b"".join(open(p, "rb").read() for p in glob.glob(str(db) + "-wal"))
    assert b"texto confidencial" not in raw and b"pregunta privada" not in raw
    with store.connect(db) as conn:
        vector_blob = conn.execute("SELECT vec FROM doc_embeddings").fetchone()["vec"]
    assert vector_blob.startswith(vault.MAGIC)
    assert np.allclose(store.load_vectors(db, "sha", "v2", "emb", 1), vectores)
    assert store.load_chunks(db, "sha", "v2")[0][3] == "texto confidencial"
    assert store.recent_messages(db, 1, 1)[0]["content"] == "pregunta privada"
    store.save_summary(db, 1, "s1", "resumen"); assert store.get_summary(db, 1) == ("s1", "resumen")


# ---------------- selección de modelo
def test_seleccion_modelo():
    inst = {"llama3.2:1b", "llama3.2:latest"}
    assert model_select.elegir(inst, 8, 0).modelo == "llama3.2:1b"
    assert model_select.elegir(inst, 16, 0).modelo == "llama3.2"
    assert model_select.elegir(inst, 32, 8).modelo == "llama3.2"          # 8b no instalado -> baja de nivel
    assert model_select.elegir({"llama3.1:8b"}, 32, 8).modelo == "llama3.1:8b"
    assert model_select.elegir(set(), 32, 8).modelo == "llama3.2:1b"       # nada instalado -> modelo mínimo
    assert model_select.elegir(inst, 4, 0, forzado="llama3.2").modelo == "llama3.2"


# ---------------- etapas
class FakeIdx:
    def __init__(self, n): self.chunks = [Chunk("c.pdf", 1 + i, i, f"Cliente {i} vive en X{i}. " * 30, None) for i in range(n)]


def test_totalidad_y_resumen(monkeypatch):
    assert rag_stages.es_totalidad("Lista todos los clientes") and rag_stages.es_totalidad("¿Cuántos clientes hay?")
    assert not rag_stages.es_totalidad("¿Cuál es la renta mensual?")
    monkeypatch.setattr(rag_stages, "_llamar", lambda m, s, u, o=None, h=None: "resumen")
    assert rag_stages.resumir_indice(FakeIdx(30), "m") == "resumen"


def test_respuesta_por_etapas(monkeypatch):
    llamadas = []
    def fake(m, s, u, o=None, h=None):
        llamadas.append(s[:10])
        return "NADA" if "Cliente 0 " not in u and "Cliente 1 " not in u else "Cliente 0\nCliente 1\nCliente 0"
    monkeypatch.setattr(rag_stages, "_llamar", fake)
    class Cli:
        def chat(self, **kw): return iter([{"message": {"content": "Son "}}, {"message": {"content": "2."}}])
    monkeypatch.setattr(rag_core, "cliente_ollama", lambda host=None: Cli())
    ev = list(rag_stages.responder_por_etapas("lista todos", FakeIdx(10), "m", {}, 4500))
    tipos = [e["type"] for e in ev]
    assert tipos.count("progress") >= 2 and "".join(e["text"] for e in ev if e["type"] == "delta") == "Son 2."
    assert ev[-1]["type"] == "sources" and ev[-1]["sources"]


# ---------------- índice parcial y nombres
def test_parcial_y_nombres(tmp_path, monkeypatch):
    monkeypatch.setattr(rag_hybrid, "_embed", lambda t, m, host=None: None)
    rutas, nombres = [], {}
    for i in range(3):
        p = tmp_path / f"{i}.txt"; vault.write_encrypted(p, f"Documento {i} habla de arriendo".encode())
        rutas.append(str(p)); nombres[str(p)] = "contrato.txt"          # mismo nombre original
    vistos = []
    idx = rag_hybrid.construir_indice_hibrido(rutas, nombres=nombres, parcial=lambda i, h, t: vistos.append((len(i.chunks), h, t)))
    assert vistos == [(1, 1, 3), (2, 2, 3)]
    assert sorted({c.doc_name for c in idx.chunks}) == ["contrato.txt", "contrato.txt (2)", "contrato.txt (3)"]


# ---------------- carpeta vigilada
def test_carpeta_vigilada(tmp_path):
    (tmp_path / "7").mkdir()
    (tmp_path / "7" / "a.txt").write_text("hola"); (tmp_path / "7" / "b.exe").write_text("x")
    (tmp_path / "7" / "~$tmp.docx").write_text("x")
    recibidos = []
    w = watcher.CarpetaVigilada(tmp_path, lambda cid, p: recibidos.append((cid, p.name)) or True, {".txt", ".docx"})
    assert w.escanear_una_vez() == []                       # primera vista: espera a que termine de copiarse
    assert len(w.escanear_una_vez()) == 1 and recibidos == [(7, "a.txt")]
    archivados = list((tmp_path / "7" / "procesados").iterdir())
    assert len(archivados) == 1 and archivados[0].read_bytes().startswith(vault.MAGIC)
    assert vault.read_bytes(archivados[0]) == b"hola"
    assert (tmp_path / "7" / "b.exe").exists()
    w.detener()
