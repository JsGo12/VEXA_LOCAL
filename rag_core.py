"""
Vexa Local — Motor RAG (núcleo, sin interfaz)
================================================

Este módulo contiene TODA la lógica de negocio del sistema RAG: extracción
de documentos, chunking inteligente, índice BM25, búsqueda de contexto y
consulta al LLM local vía Ollama. No depende de ninguna librería de
interfaz (ni Streamlit ni Tkinter), para poder reutilizarse desde
cualquier capa visual — en este proyecto, la app de escritorio
`desktop_app.py`.

Todo el procesamiento ocurre en la máquina local. No se realiza ninguna
llamada a servicios externos: solo se habla con el servidor local de
Ollama (http://localhost:11434).
"""

from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from typing import List, Dict, Any, Optional
import os

import pdfplumber
import ollama
from rank_bm25 import BM25Okapi
from docx import Document as DocxDocument


# =============================================================================
# CONFIGURACIÓN GLOBAL
# =============================================================================

APP_TITLE = "Vexa Local"
APP_SUBTITLE = "Chat con documentos 100% offline — sin nube, sin fugas de datos"

# Modelos livianos recomendados para correr en CPU/GPU modesta vía Ollama.
MODELOS_DISPONIBLES = ["llama3.2", "phi3", "llama3.2:1b", "phi3:mini"]

# Tamaño de los chunks de texto (en palabras) y solapamiento entre ellos.
CHUNK_SIZE_PALABRAS = 220
CHUNK_OVERLAP_PALABRAS = 40

# Cantidad de fragmentos de contexto que se recuperan por pregunta puntual.
TOP_K_FRAGMENTOS = 3

# Cantidad de fragmentos cuando la pregunta pide un LISTADO/TOTALIDAD de
# registros (ej. "todos los pacientes de X"), en vez de un registro puntual.
TOP_K_FRAGMENTOS_AGREGACION = 8

# Host del servidor local de Ollama (por defecto ya apunta a localhost).
OLLAMA_HOST = "http://localhost:11434"


# =============================================================================
# MODELOS DE DATOS
# =============================================================================

@dataclass
class Chunk:
    """Representa un fragmento de texto extraído de un documento,
    conservando la metadata necesaria para poder citar la fuente."""

    doc_name: str
    page_number: int
    chunk_index: int
    text: str
    # Página final del fragmento, solo si un registro cruza un salto de
    # página del PDF (ej. una ficha de "Cliente N" partida entre dos
    # páginas). Si es None, se asume que el chunk vive en una sola página.
    page_number_fin: Optional[int] = None

    @property
    def fuente(self) -> str:
        """Etiqueta legible para citar este fragmento como fuente."""
        fin = self.page_number_fin or self.page_number
        if fin != self.page_number:
            return f"{self.doc_name} — páginas {self.page_number}-{fin}"
        return f"{self.doc_name} — página {self.page_number}"


@dataclass
class IndiceDocumental:
    """Contenedor del índice BM25 junto con los chunks originales."""

    chunks: List[Chunk] = field(default_factory=list)
    bm25: Optional[BM25Okapi] = None

    def esta_vacio(self) -> bool:
        return len(self.chunks) == 0


# =============================================================================
# MÓDULO DE PROCESAMIENTO DE DOCUMENTOS
# =============================================================================

def _tokenizar(texto: str) -> List[str]:
    """Tokeniza texto para BM25: minúsculas, sin tildes, solo palabras.

    BM25 funciona por coincidencia léxica exacta, por lo que normalizar
    (quitar tildes/mayúsculas) mejora el recall en español sin necesidad
    de embeddings.
    """
    texto = texto.lower()
    texto = unicodedata.normalize("NFKD", texto)
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    tokens = re.findall(r"[a-z0-9]+", texto)
    return tokens


def extraer_texto_por_pagina(ruta_o_archivo) -> List[Dict[str, Any]]:
    """Extrae el texto de un PDF, página por página.

    Args:
        ruta_o_archivo: ruta de archivo (str/Path) o un objeto tipo
            archivo ya abierto (pdfplumber acepta ambos).

    Returns:
        Lista de dicts {"page_number": int, "text": str}, uno por página
        que contenga texto extraíble.
    """
    paginas = []
    with pdfplumber.open(ruta_o_archivo) as pdf:
        for i, pagina in enumerate(pdf.pages, start=1):
            texto = pagina.extract_text() or ""
            texto = texto.strip()
            if texto:
                paginas.append({"page_number": i, "text": texto})
    return paginas


def extraer_texto_docx(
    ruta_o_archivo, parrafos_por_pagina: int = 25
) -> List[Dict[str, Any]]:
    """Extrae el texto de un archivo .docx, simulando 'páginas' agrupando
    bloques de contenido, ya que Word no guarda paginación real en su
    estructura interna.

    IMPORTANTE: recorre el documento en su orden real (párrafos y tablas
    intercalados tal como aparecen), incluyendo el texto de celdas de
    tablas — `doc.paragraphs` por sí solo NO incluye el contenido de
    tablas, así que fichas con formato de tabla (ej. "Cliente N",
    "Paciente N") quedarían casi vacías si solo se leyeran los párrafos.

    Args:
        ruta_o_archivo: ruta de archivo (str/Path) o un objeto tipo
            archivo ya abierto.
        parrafos_por_pagina: cantidad de bloques que forman cada "página"
            simulada, usada solo para poder citar una referencia aproximada.

    Returns:
        Lista de dicts {"page_number": int, "text": str}.
    """
    doc = DocxDocument(ruta_o_archivo)
    bloques: List[str] = []

    def _texto_de_tabla(tabla) -> str:
        """Convierte una tabla en texto plano, una fila por línea,
        con las celdas separadas por ' | ' para preservar la relación
        campo-valor típica de las fichas."""
        lineas = []
        for fila in tabla.rows:
            celdas = [c.text.strip() for c in fila.cells]
            celdas = [c for c in celdas if c]
            if celdas:
                lineas.append(" | ".join(celdas))
        return "\n".join(lineas)

    # Recorre el XML del cuerpo del documento en orden, para intercalar
    # párrafos y tablas tal como aparecen visualmente en el Word.
    cuerpo = doc.element.body
    mapa_parrafos = {p._element: p for p in doc.paragraphs}
    mapa_tablas = {t._element: t for t in doc.tables}

    for elemento in cuerpo.iterchildren():
        if elemento in mapa_parrafos:
            texto = mapa_parrafos[elemento].text.strip()
            if texto:
                bloques.append(texto)
        elif elemento in mapa_tablas:
            texto_tabla = _texto_de_tabla(mapa_tablas[elemento])
            if texto_tabla:
                bloques.append(texto_tabla)

    paginas = []
    for i in range(0, len(bloques), parrafos_por_pagina):
        grupo = bloques[i: i + parrafos_por_pagina]
        texto = "\n".join(grupo)
        if texto.strip():
            paginas.append(
                {"page_number": (i // parrafos_por_pagina) + 1, "text": texto}
            )
    return paginas


def _detectar_encabezado_registro(texto_completo: str) -> Optional[re.Pattern]:
    """Detecta AUTOMÁTICAMENTE si el documento está estructurado en fichas
    repetidas (ej. "Cliente 1", "Paciente 2", "Ticket 003", "Empleado N",
    "Producto 45"...) sin asumir ninguna palabra fija de antemano — cada
    empresa nombra sus registros de forma distinta.

    Busca líneas cortas (1-2 palabras) que terminen en un número, y que
    esa MISMA etiqueta se repita al menos 2 veces con números distintos
    a lo largo del documento. Eso es un indicio fuerte de que se trata
    de una serie de fichas/registros, no de texto corrido.

    Returns:
        Un patrón compilado (para usar con re.split, con lookahead) que
        coincide con el inicio de cada ficha detectada; o None si el
        documento no muestra ese patrón (texto corrido normal).
    """
    patron_linea = re.compile(
        r"^[ \t]*([A-ZÁÉÍÓÚÑ][\wÁÉÍÓÚÑáéíóúñ]*(?:[ \t]+[A-ZÁÉÍÓÚÑ][\wÁÉÍÓÚÑáéíóúñ]*)?)"
        r"[ \t]*[:#°ºNn.]*[ \t]*(\d{1,6})[ \t]*$",
        re.MULTILINE,
    )

    conteo_por_etiqueta: Dict[str, set] = {}
    for match in patron_linea.finditer(texto_completo):
        etiqueta = match.group(1).strip().lower()
        numero = match.group(2)
        conteo_por_etiqueta.setdefault(etiqueta, set()).add(numero)

    # Solo cuentan etiquetas con al menos 2 apariciones con NÚMEROS
    # DISTINTOS (una sola vez podría ser una fecha, un total, etc.,
    # no necesariamente el separador real de fichas).
    candidatos = {
        etiqueta: numeros
        for etiqueta, numeros in conteo_por_etiqueta.items()
        if len(numeros) >= 2
    }

    if not candidatos:
        return None

    # Se elige la etiqueta con más repeticiones: la más probable de ser
    # el separador real de fichas del documento.
    mejor_etiqueta = max(candidatos, key=lambda e: len(candidatos[e]))

    return re.compile(
        r"(?=^[ \t]*" + re.escape(mejor_etiqueta) + r"[ \t]*[:#°ºNn.]*[ \t]*\d+\b)",
        re.IGNORECASE | re.MULTILINE,
    )


def dividir_en_chunks(
    doc_name: str,
    paginas: List[Dict[str, Any]],
    chunk_size: int = CHUNK_SIZE_PALABRAS,
    overlap: int = CHUNK_OVERLAP_PALABRAS,
) -> List[Chunk]:
    """Divide el texto del documento en chunks, preservando la página (o
    rango de páginas) como metadata para poder citarla luego.

    Estrategia:
    1. Se detecta AUTOMÁTICAMENTE (sin palabras fijas) si el documento
       está estructurado en fichas repetidas — "Cliente 1", "Paciente 2",
       "Ticket 003", "Empleado N", lo que sea que use cada empresa. Si es
       así, cada ficha completa se convierte en UN SOLO chunk atómico —
       incluso si se extiende a través de un salto de página del PDF.
       Esto evita que el corte por palabras (o el borde de página) parta
       una ficha a la mitad, o que se mezclen campos de dos fichas
       distintas al momento de la búsqueda.
    2. Si no se detecta ese patrón (documento de texto corrido normal),
       se usa el chunking genérico por cantidad de palabras con solape,
       sin cruzar páginas.
    """
    if not paginas:
        return []

    texto_completo_doc = "\n".join(p["text"] for p in paginas)
    patron_split = _detectar_encabezado_registro(texto_completo_doc)

    if patron_split is not None:
        return _dividir_por_registros(doc_name, paginas, patron_split)

    return _dividir_por_palabras(doc_name, paginas, chunk_size, overlap)


def _dividir_por_registros(
    doc_name: str, paginas: List[Dict[str, Any]], patron_split: re.Pattern
) -> List[Chunk]:
    """Chunking para documentos tipo 'ficha' (Cliente 1, Paciente 2, ...).

    Concatena todas las páginas insertando un marcador invisible entre
    ellas para poder rastrear en qué página(s) cae cada registro, y luego
    divide el texto completo usando el patrón de ficha ya detectado
    (`patron_split`, específico de este documento). Así, un registro que
    el PDF/Word cortó justo en un salto de página queda igualmente
    completo en un solo chunk, citando el rango real (ej. "páginas 1-2").
    """
    marcador = "\x00PAGE:{}\x00"
    partes = []
    for pagina in paginas:
        partes.append(marcador.format(pagina["page_number"]))
        partes.append(pagina["text"])
    texto_completo = "\n".join(partes)

    patron_marcador = re.compile(r"\x00PAGE:(\d+)\x00")

    bloques = patron_split.split(texto_completo)

    chunks: List[Chunk] = []
    pagina_actual = paginas[0]["page_number"]
    idx_chunk = 0

    for bloque in bloques:
        paginas_en_bloque = [int(p) for p in patron_marcador.findall(bloque)]
        texto_limpio = patron_marcador.sub("", bloque).strip()

        # La página de INICIO del bloque es la que estaba vigente justo
        # antes de entrar a este bloque (heredada del bloque anterior),
        # no la del primer marcador que aparezca dentro de él — de lo
        # contrario, un registro que empieza en la página 1 y solo cruza
        # a la página 2 a mitad de camino quedaría mal etiquetado como
        # si empezara en la página 2.
        pagina_inicio = pagina_actual

        if paginas_en_bloque:
            pagina_fin = paginas_en_bloque[-1]
            pagina_actual = pagina_fin
        else:
            pagina_fin = pagina_actual

        if not texto_limpio:
            continue

        chunks.append(
            Chunk(
                doc_name=doc_name,
                page_number=pagina_inicio,
                page_number_fin=pagina_fin,
                chunk_index=idx_chunk,
                text=texto_limpio,
            )
        )
        idx_chunk += 1

    return chunks


def _dividir_por_palabras(
    doc_name: str,
    paginas: List[Dict[str, Any]],
    chunk_size: int,
    overlap: int,
) -> List[Chunk]:
    """Chunking genérico por cantidad de palabras con solape, usado para
    documentos de texto corrido (no fichas). No cruza páginas."""
    chunks: List[Chunk] = []

    for pagina in paginas:
        palabras = pagina["text"].split()
        if not palabras:
            continue

        paso = max(chunk_size - overlap, 1)
        idx_chunk_local = 0

        for inicio in range(0, len(palabras), paso):
            fragmento_palabras = palabras[inicio: inicio + chunk_size]
            if not fragmento_palabras:
                continue

            texto_chunk = " ".join(fragmento_palabras)
            chunks.append(
                Chunk(
                    doc_name=doc_name,
                    page_number=pagina["page_number"],
                    chunk_index=idx_chunk_local,
                    text=texto_chunk,
                )
            )
            idx_chunk_local += 1

            # Si ya cubrimos toda la página, no seguimos iterando.
            if inicio + chunk_size >= len(palabras):
                break

    return chunks


def construir_indice(chunks: List[Chunk]) -> IndiceDocumental:
    """Construye el índice BM25 a partir de la lista de chunks."""
    if not chunks:
        return IndiceDocumental(chunks=[], bm25=None)

    corpus_tokenizado = [_tokenizar(c.text) for c in chunks]
    bm25 = BM25Okapi(corpus_tokenizado)
    return IndiceDocumental(chunks=chunks, bm25=bm25)


def procesar_documentos(rutas_archivos: List[str]) -> IndiceDocumental:
    """Pipeline completo: extrae texto, genera chunks y construye el
    índice BM25 para uno o varios documentos (PDF o Word .docx,
    detectado automáticamente por la extensión del archivo).

    Args:
        rutas_archivos: lista de rutas de archivo en disco (str).
    """
    import os

    todos_los_chunks: List[Chunk] = []

    for ruta in rutas_archivos:
        nombre_doc = os.path.basename(ruta)

        if ruta.lower().endswith(".docx"):
            paginas = extraer_texto_docx(ruta)
        else:
            paginas = extraer_texto_por_pagina(ruta)

        chunks_doc = dividir_en_chunks(doc_name=nombre_doc, paginas=paginas)
        todos_los_chunks.extend(chunks_doc)

    return construir_indice(todos_los_chunks)


# =============================================================================
# MÓDULO DE BÚSQUEDA (RETRIEVAL)
# =============================================================================

_PALABRAS_CLAVE_AGREGACION = (
    "todos", "todas", "cada", "lista", "listado", "listame", "lístame",
    "cuales", "cuáles", "cuantos", "cuántos", "cuantas", "cuántas",
    "enumera", "enumerame", "enumérame",
)


def _es_pregunta_de_listado(pregunta: str) -> bool:
    """Detecta si la pregunta pide un listado/totalidad de registros
    (en vez de apuntar a un registro puntual), para poder ampliar la
    cantidad de fragmentos recuperados y no perder resultados relevantes
    frente a coincidencias léxicas ruidosas (ej. el título del documento)."""
    tokens = set(_tokenizar(pregunta))
    return any(palabra in tokens for palabra in _PALABRAS_CLAVE_AGREGACION)


def determinar_top_k(pregunta: str, total_chunks: int) -> int:
    """Decide cuántos fragmentos pedir a BM25 según el tipo de pregunta.

    Si la pregunta parece pedir un listado o totalidad de registros
    ("todos los pacientes de X", "cuáles clientes tienen Y"), se usa un
    top_k más generoso para no perder registros relevantes cuando hay
    fragmentos ruidosos (como el título del documento) compitiendo por
    los primeros lugares del ranking. Nunca se pide más fragmentos que
    los que existen en el índice.
    """
    if _es_pregunta_de_listado(pregunta):
        return min(total_chunks, TOP_K_FRAGMENTOS_AGREGACION)
    return min(total_chunks, TOP_K_FRAGMENTOS)


def buscar_contexto_relevante(
    pregunta: str,
    indice: IndiceDocumental,
    top_k: int = TOP_K_FRAGMENTOS,
) -> List[Chunk]:
    """Busca los `top_k` chunks más relevantes para la pregunta usando BM25.

    Args:
        pregunta: pregunta del usuario en lenguaje natural.
        indice: índice documental ya construido.
        top_k: cantidad de fragmentos a recuperar.

    Returns:
        Lista de Chunks ordenados de mayor a menor relevancia. Puede
        devolver una lista vacía si el índice está vacío o no hay
        coincidencias con score positivo.
    """
    if indice.esta_vacio() or indice.bm25 is None:
        return []

    tokens_pregunta = _tokenizar(pregunta)
    if not tokens_pregunta:
        return []

    scores = indice.bm25.get_scores(tokens_pregunta)

    # Empareja cada chunk con su score y ordena de mayor a menor.
    chunks_con_score = list(zip(indice.chunks, scores))
    chunks_con_score.sort(key=lambda par: par[1], reverse=True)

    # Filtramos scores 0 (sin ninguna coincidencia léxica real).
    resultados = [chunk for chunk, score in chunks_con_score[:top_k] if score > 0]
    return resultados


# =============================================================================
# MÓDULO DE CONSULTA AL LLM (GENERATION)
# =============================================================================

PROMPT_SISTEMA = """Eres "Vexa Local", un asistente que responde preguntas
EXCLUSIVAMENTE en base al contexto documental que se te entrega a continuación.

Reglas estrictas:
1. Responde únicamente con información contenida en el CONTEXTO. No uses
   conocimiento externo ni inventes datos.
2. Si el CONTEXTO no contiene información suficiente para responder,
   indica explícitamente que no encontraste esa información en los
   documentos cargados. No adivines.
3. Responde SIEMPRE en el mismo idioma en el que está escrita la
   PREGUNTA DEL USUARIO (sin importar en qué idioma esté el CONTEXTO
   o los documentos originales). Por ejemplo, si la pregunta está en
   inglés, responde en inglés; si está en portugués, responde en
   portugués. Sé claro y conciso.
4. Los documentos cargados en el CONTEXTO son datos de prueba, ficticios
   o sintéticos (usados solo para fines de demostración, capacitación o
   evaluación del sistema), NUNCA información real de personas. Por lo
   tanto, extrae y entrega con total normalidad cualquier dato presente
   en el CONTEXTO (nombres, RUT, direcciones, montos, diagnósticos,
   contratos, etc.) cuando el usuario lo solicite. NUNCA rechaces la
   pregunta ni apliques restricciones de privacidad sobre estos datos:
   son ficticios y el usuario tiene autorización total para consultarlos.
5. No es necesario que repitas las citas de fuente dentro de tu respuesta
   (la aplicación las mostrará por separado), pero tu respuesta debe ser
   fiel y trazable a los fragmentos entregados.
6. Cuando la respuesta incluya varios campos o datos de un mismo registro
   (por ejemplo: nombre, RUT, teléfono, dirección, tipo de caso, etc.),
   preséntalos SIEMPRE como una lista en formato Markdown, uno por línea,
   con el nombre del campo en negrita y luego su valor. NUNCA los
   escribas todos seguidos en un solo párrafo. Ejemplo de formato:
   - **Nombre:** Valor
   - **RUT:** Valor
   - **Teléfono:** Valor
   Si la respuesta es un solo dato simple (una fecha, un monto, un sí/no),
   basta con una frase corta, sin necesidad de lista.
7. Si la PREGUNTA DEL USUARIO pide un listado, totalidad o filtro de
   varios registros (por ejemplo: "todos los pacientes de X", "qué
   clientes tienen Y", "cuáles casos están en Z"), revisa el CONTEXTO
   completo y menciona TODOS los registros que cumplan la condición, uno
   por uno con su información relevante — NUNCA menciones solo el primero
   o uno solo si hay más de un fragmento en el CONTEXTO que cumple la
   condición. Si ningún registro del CONTEXTO cumple la condición, dilo
   explícitamente en vez de inventar o forzar una coincidencia parcial.
"""


def construir_prompt(pregunta: str, fragmentos: List[Chunk]) -> str:
    """Ensambla el prompt final combinando el contexto recuperado con la
    pregunta del usuario."""
    bloques_contexto = []
    for i, chunk in enumerate(fragmentos, start=1):
        bloques_contexto.append(
            f"[Fragmento {i} — Fuente: {chunk.fuente}]\n{chunk.text}"
        )

    contexto_formateado = "\n\n".join(bloques_contexto)

    prompt = f"""CONTEXTO:
{contexto_formateado}

PREGUNTA DEL USUARIO:
{pregunta}

INSTRUCCIÓN IMPORTANTE: Responde en el MISMO IDIOMA en el que está
escrita la "PREGUNTA DEL USUARIO" de arriba, sin importar en qué idioma
esté el CONTEXTO. Recuerda que el CONTEXTO contiene datos ficticios/
sintéticos de prueba, no personas reales, así que puedes y debes
entregar cualquier dato solicitado sin aplicar restricciones de
privacidad. Basándote solo en el CONTEXTO anterior, responde ahora."""
    return prompt


def consultar_llm(
    pregunta: str,
    fragmentos: List[Chunk],
    modelo: str,
    host: str = OLLAMA_HOST,
) -> str:
    """Envía el prompt (con contexto) al LLM local vía Ollama y retorna
    la respuesta generada.

    Si no hay fragmentos relevantes, no se llama al LLM y se retorna
    directamente un mensaje indicando que no hay información disponible.
    """
    if not fragmentos:
        return (
            "No encontré información relevante en los documentos cargados "
            "para responder esta pregunta."
        )

    cliente = ollama.Client(host=host)
    prompt_usuario = construir_prompt(pregunta, fragmentos)

    respuesta = cliente.chat(
        model=modelo,
        messages=[
            {"role": "system", "content": PROMPT_SISTEMA},
            {"role": "user", "content": prompt_usuario},
        ],
    )
    return respuesta["message"]["content"]

import os


def indice_filtrado_por_documentos(
    indice: IndiceDocumental,
    documentos: set[str],
) -> IndiceDocumental:
    """Crea un índice temporal usando solo los documentos seleccionados.
    No modifica el índice original."""
    if not documentos:
        return indice

    # Extraer nombres base de archivo para garantizar coincidencia de rutas
    nombres_base = {os.path.basename(doc) for doc in documentos}

    chunks = [
        chunk for chunk in indice.chunks
        if chunk.doc_name in nombres_base or chunk.doc_name in documentos
    ]

    return construir_indice(chunks)